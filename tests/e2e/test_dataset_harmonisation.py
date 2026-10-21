"""End-to-end test for harmonising two differently-shaped spreadsheets into a
shared set of canonical Fields via the Datasets standardisation wizard.

Scenario (real data, test-data/dataset_std1_fixed.csv + dataset_std2.xlsx):
  1. Standardise dataset_std1_fixed.csv, establishing canonical Fields
     (Date, Cultivar, FruitHeight [H/L], FruitNumber, Diameter, Weight).
  2. Standardise dataset_std2.xlsx's "RawData" sheet, whose columns are named
     and cased differently (Height instead of FruitHeight, lowercase
     cultivar names, Rep instead of FruitNumber) - harmonising it into the
     same Fields. This adds "Low"/"High" as new FruitHeight values alongside
     the existing H/L, and creates new Rep/HarvestDate/AssessmentDate Fields.
  3. Retire the old H/L values and the now-superseded FruitNumber/Date
     Fields, completing the migration to the new canonical names.
  4. Re-run both datasets' standardisation: dataset_std1_fixed now needs its
     Date/FruitNumber columns remapped to the existing HarvestDate/Rep
     Fields, and its raw H/L values remapped to the existing High/Low
     values. dataset_std2_RawData's own columns/values now match the new
     canonical names automatically, needing only the same Height->FruitHeight
     field-level fix as before.
  5. Download both datasets' standardised output and verify the RawData
     sheet reflects the fully harmonised schema.

Two real (not app-bug) quirks confirmed by exploring this live, both handled
below: difflib's fuzzy header matching sometimes proposes a wrong-but-similar
existing Field ("Height" naively best-matches "Weight" via a shared "eight"
suffix, "Date" naively best-matches "Diameter") - the switch is already
correctly on "Existing" in these cases, only the dropdown selection itself
needs correcting.
"""

import openpyxl
import pytest
from playwright.sync_api import expect

from tests.pages.dataset_standardise_page import DatasetStandardisePage
from tests.pages.datasets_page import DatasetsPage
from tests.pages.field_detail_page import FieldDetailPage
from tests.pages.fields_page import FieldsPage
from utils.data_loader import data_file_path


def open_standardise(datasets_page: DatasetsPage, name: str) -> DatasetStandardisePage:
    datasets_page.go_to_standardise_wizard(name)
    return DatasetStandardisePage(datasets_page.page)


def go_to_fields(datasets_page: DatasetsPage) -> None:
    datasets_page.page.get_by_role("link", name="Fields", exact=True).click()


def go_to_datasets(datasets_page: DatasetsPage) -> None:
    datasets_page.page.get_by_role("link", name="Datasets", exact=True).click()


@pytest.mark.regression
def test_harmonise_two_datasets_into_shared_fields(datasets_page: DatasetsPage, tmp_path):
    # --- Round 1: dataset_std1_fixed.csv establishes the canonical Fields ---
    datasets_page.upload(data_file_path("dataset_std1_fixed.csv"))

    standardise = open_standardise(datasets_page, "dataset_std1_fixed")
    # Columns in order: Date, Cultivar, FruitHeight, FruitNumber, Diameter, Weight.
    standardise.select_type(3, "Number")  # FruitNumber defaults to Measurement.
    standardise.go_to_page_2()
    # Accept defaults: Cultivar (Gala/Envy) and FruitHeight (H/L) both become
    # new Category values under their respective new Fields.
    standardise.finish()

    expect(datasets_page.status_badge()).to_have_text("Complete")

    # --- Round 2: dataset_std2.xlsx (RawData sheet) harmonises into the same Fields ---
    datasets_page.upload_choosing_sheet(data_file_path("dataset_std2.xlsx"), "RawData")

    standardise = open_standardise(datasets_page, "dataset_std2_RawData")
    # Columns in order: HarvestDate, AssessmentDate, Cultivar, Height, Rep, Diameter, Weight.
    # Height already defaults to "Existing" but wrongly guesses "Weight".
    standardise.select_existing_field(3, "FruitHeight")
    # Rep defaults to a new Measurement field - override the type to Number.
    standardise.select_type_with_existing_fields(4, "Number")
    standardise.go_to_page_2()
    # Accept defaults: Cultivar's "envy"/"gala" already auto-match existing
    # Envy/Gala; Height's "Low"/"High" default to new values, added
    # alongside the existing H/L.
    standardise.finish()

    expect(datasets_page.status_badge()).to_have_text("Complete")

    # --- Retire the old H/L values and the now-superseded Fields ---
    go_to_fields(datasets_page)
    fields_page = FieldsPage(datasets_page.page)
    fields_page.open_field("FruitHeight")
    field_detail = FieldDetailPage(datasets_page.page)
    field_detail.delete_first_value_row()  # H
    field_detail.delete_first_value_row()  # L
    field_detail.save()
    expect(field_detail.banner).to_contain_text("Field FruitHeight saved successfully.")

    go_to_fields(datasets_page)
    fields_page.delete_field("FruitNumber").confirm()
    fields_page.delete_field("Date").confirm()

    # --- Round 3: re-run dataset_std1_fixed, remapping onto the new canonical Fields ---
    go_to_datasets(datasets_page)
    standardise = open_standardise(datasets_page, "dataset_std1_fixed")
    # Date already defaults to "Existing" but wrongly guesses "Diameter".
    standardise.select_existing_field(0, "HarvestDate")
    # FruitNumber's old target Field is gone - it now defaults to "New".
    standardise.select_existing_switch(3)
    standardise.select_existing_field(3, "Rep")
    standardise.go_to_page_2()
    # Value order on this page: Cultivar's Envy(0)/Gala(1) already correctly
    # auto-match existing values; FruitHeight's H(2)/L(3) default to "New"
    # (no longer close enough to the remaining High/Low) - fix both.
    standardise.select_value_existing_switch(2)
    standardise.select_existing_value(2, "High")
    standardise.select_value_existing_switch(3)
    standardise.select_existing_value(3, "Low")
    standardise.finish()

    expect(datasets_page.status_badge()).to_have_text("Complete")

    download = datasets_page.download_standardised("dataset_std1_fixed")
    std1_path = tmp_path / "std1_standardised.xlsx"
    download.save_as(std1_path)

    wb = openpyxl.load_workbook(std1_path, data_only=True)
    rows = list(wb["RawData"].iter_rows(values_only=True))
    assert rows[0] == ("HarvestDate", "Cultivar", "FruitHeight", "Rep", "Diameter", "Weight")
    assert rows[1:] == [
        ("2025-06-01", "Gala", "Low", 1, 11.5, 0.2),
        ("2025-06-01", "Gala", "Low", 2, 12.1, 0.21),
        ("2025-06-01", "Gala", "Low", 3, 11, 0.19),
        ("2025-06-01", "Gala", "High", 1, 12.3, 0.21),
        ("2025-06-01", "Gala", "High", 2, 13.6, 0.22),
        ("2025-06-01", "Gala", "High", 3, 11.3, 0.2),
        ("2025-06-02", "Envy", "Low", 1, 10.7, 0.2),
        ("2025-06-02", "Envy", "Low", 2, 12.8, 0.21),
        ("2025-06-02", "Envy", "Low", 3, 10.4, 0.19),
        ("2025-06-02", "Envy", "High", 1, 12.1, 0.22),
        ("2025-06-02", "Envy", "High", 2, 12.8, 0.24),
        ("2025-06-02", "Envy", "High", 3, 11.9, 0.2),
    ]

    # --- Round 4: re-run dataset_std2_RawData - now matches the canonical names automatically ---
    standardise = open_standardise(datasets_page, "dataset_std2_RawData")
    # Height still needs the same field-level fix as round 2 (the wrong
    # "Weight" guess doesn't change just because FruitHeight's values did).
    standardise.select_existing_field(3, "FruitHeight")
    standardise.go_to_page_2()
    # Accept defaults: both Cultivar's and Height's values now auto-match
    # the existing Fields' current values exactly.
    standardise.finish()

    expect(datasets_page.status_badge()).to_have_text("Complete")

    download2 = datasets_page.download_standardised("dataset_std2_RawData")
    std2_path = tmp_path / "std2_standardised.xlsx"
    download2.save_as(std2_path)

    wb2 = openpyxl.load_workbook(std2_path, data_only=True)
    rows2 = list(wb2["RawData"].iter_rows(values_only=True))
    assert rows2[0] == (
        "HarvestDate", "AssessmentDate", "Cultivar", "FruitHeight", "Rep", "Diameter", "Weight",
    )
    assert rows2[1:] == [
        ("2026-06-01", "2026-06-03", "Gala", "Low", 1, 12.4, 0.24),
        ("2026-06-01", "2026-06-03", "Gala", "Low", 2, 11.1, 0.2),
        ("2026-06-01", "2026-06-03", "Gala", "Low", 3, 11.9, 0.19),
        ("2026-06-01", "2026-06-03", "Gala", "High", 1, 12.3, 0.21),
        ("2026-06-01", "2026-06-03", "Gala", "High", 2, 10.7, 0.25),
        ("2026-06-01", "2026-06-03", "Gala", "High", 3, 13.3, 0.23),
        ("2026-06-02", "2026-06-03", "Envy", "Low", 1, 13.7, None),
        ("2026-06-02", "2026-06-03", "Envy", "Low", 2, 12.1, 0.22),
        ("2026-06-02", "2026-06-03", "Envy", "Low", 3, 11.8, 0.19),
        ("2026-06-02", "2026-06-03", "Envy", "High", 1, 10.4, 0.22),
        ("2026-06-02", "2026-06-03", "Envy", "High", 2, 12.1, 0.23),
        ("2026-06-02", "2026-06-03", "Envy", "High", 3, 11.9, 0.2),
    ]
