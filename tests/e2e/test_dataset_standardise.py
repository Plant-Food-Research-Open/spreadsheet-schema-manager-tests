import openpyxl
import pytest
from playwright.sync_api import expect

from tests.pages.base_page import click_until, pressable
from tests.pages.dataset_standardise_page import DatasetStandardisePage
from tests.pages.datasets_page import DatasetsPage
from tests.pages.fields_page import FieldsPage
from utils.data_loader import data_file_path


def write_file(tmp_path, filename: str, content: bytes) -> str:
    path = tmp_path / filename
    path.write_bytes(content)
    return str(path)


def open_standardise(datasets_page: DatasetsPage, name: str) -> DatasetStandardisePage:
    datasets_page.go_to_standardise_wizard(name)
    return DatasetStandardisePage(datasets_page.page)


@pytest.mark.regression
def test_standardise_page1_loads(datasets_page: DatasetsPage):
    datasets_page.upload(data_file_path("dataset_ready.csv"))

    standardise = open_standardise(datasets_page, "dataset_ready")

    expect(standardise.heading("dataset_ready")).to_be_visible()
    expect(standardise.next_button).to_be_visible()


@pytest.mark.regression
def test_standardise_happy_path_marks_dataset_complete(datasets_page: DatasetsPage):
    datasets_page.upload(data_file_path("dataset_ready.csv"))

    standardise = open_standardise(datasets_page, "dataset_ready")
    standardise.go_to_page_2()
    standardise.finish()

    expect(datasets_page.status_badge()).to_have_text("Complete")
    expect(datasets_page.view_standardised_button("dataset_ready")).to_be_enabled()
    expect(datasets_page.download_standardised_button("dataset_ready")).to_be_enabled()


@pytest.mark.regression
def test_standardise_back_returns_to_page1(datasets_page: DatasetsPage):
    datasets_page.upload(data_file_path("dataset_ready.csv"))

    standardise = open_standardise(datasets_page, "dataset_ready")
    standardise.go_to_page_2()
    standardise.back_to_page_1()

    expect(standardise.next_button).to_be_visible()


@pytest.mark.regression
def test_standardisation_log_sheet_records_mapping_decisions(datasets_page: DatasetsPage, tmp_path):
    path = write_file(tmp_path, "qa_log.csv", b"Grade\nA\nB\nA\n")
    datasets_page.upload(path)

    standardise = open_standardise(datasets_page, "qa_log")
    standardise.go_to_page_2()
    standardise.finish()

    expect(datasets_page.status_badge()).to_have_text("Complete")

    download = datasets_page.download_standardised("qa_log")
    download_path = tmp_path / "qa_log_standardised.xlsx"
    download.save_as(download_path)

    wb = openpyxl.load_workbook(download_path, data_only=True)
    assert "StandardisationLog" in wb.sheetnames
    rows = list(wb["StandardisationLog"].iter_rows(values_only=True))

    assert rows[0] == (
        "Original column", "Mapped to", "Field type", "Field source",
        "Date order", "Original value", "Mapped value", "Value source",
    )
    # xlsxwriter writes blank log fields as an empty string, but openpyxl
    # reads that back as None - Excel doesn't distinguish an empty-string
    # cell from a genuinely blank one.
    assert rows[1] == ("Grade", "Grade", "Category", "New Field", None, None, None, None)
    value_rows = {row[5]: row for row in rows[2:]}
    assert value_rows["A"] == (None, None, None, None, None, "A", "A", "New Value")
    assert value_rows["B"] == (None, None, None, None, None, "B", "B", "New Value")


@pytest.mark.regression
def test_standardise_category_exceeding_max_values_shows_error(datasets_page: DatasetsPage, tmp_path):
    # 60 unique values in one column - validate_format auto-detects this as
    # "Text" (>50 uniques), so the type must be manually switched to
    # "Category" to trigger the mapping-limit check.
    rows = "\n".join(f"Value{i}" for i in range(60))
    csv_content = f"Col\n{rows}\n".encode()
    path = write_file(tmp_path, "qa_toomany.csv", csv_content)
    datasets_page.upload(path)

    standardise = open_standardise(datasets_page, "qa_toomany")
    standardise.select_type(0, "Category")
    standardise.go_to_page_2()

    expect(standardise.page.get_by_text("exceeds the mapping limit", exact=False)).to_be_visible()
    expect(pressable(standardise.finish_button)).to_have_attribute("aria-disabled", "true")


@pytest.mark.regression
def test_standardise_duplicate_field_mapping_shows_error(datasets_page: DatasetsPage, tmp_path):
    datasets_page.page.get_by_role("link", name="Fields", exact=True).click()
    FieldsPage(datasets_page.page).create_field("QA Shared Field", "Text")

    datasets_page.page.get_by_role("link", name="Datasets", exact=True).click()
    path = write_file(tmp_path, "qa_dup_map.csv", b"ColA,ColB\nx,y\n")
    datasets_page.upload(path)

    standardise = open_standardise(datasets_page, "qa_dup_map")
    standardise.select_existing_switch(0)
    standardise.select_existing_field(0, "/ Text / QA Shared Field")
    standardise.select_existing_switch(1)
    standardise.select_existing_field(1, "/ Text / QA Shared Field")
    click_until(standardise.banner, standardise.next_button, "visible")

    expect(standardise.banner).to_contain_text(
        "Multiple headers are being mapped to Field QA Shared Field."
    )


@pytest.mark.regression
def test_standardise_empty_category_field_shows_error(datasets_page: DatasetsPage, tmp_path):
    datasets_page.page.get_by_role("link", name="Fields", exact=True).click()
    FieldsPage(datasets_page.page).create_field("QA Empty Category Field", "Category")

    datasets_page.page.get_by_role("link", name="Datasets", exact=True).click()
    path = write_file(tmp_path, "qa_empty_cat.csv", b"Col\nx\n")
    datasets_page.upload(path)

    standardise = open_standardise(datasets_page, "qa_empty_cat")
    standardise.select_existing_switch(0)
    standardise.select_existing_field(0, "/ Category / QA Empty Category Field")
    click_until(standardise.banner, standardise.next_button, "visible")

    expect(standardise.banner).to_contain_text(
        "Unable to map to Field QA Empty Category Field with empty category values."
    )


@pytest.mark.regression
def test_ambiguous_date_column_defaults_dayfirst_and_can_switch_to_month_first(datasets_page: DatasetsPage, tmp_path):
    # Every row's day value is <=12, so neither a day-first nor a month-first
    # reading can be ruled out from the data alone (SSM-29 ambiguity case).
    path = write_file(tmp_path, "qa_ambiguous_date.csv", b"Col\n03/04/2024\n05/06/2024\n")
    datasets_page.upload(path)

    standardise = open_standardise(datasets_page, "qa_ambiguous_date")
    standardise.go_to_page_2()

    expect(standardise.page.get_by_text("Ambiguous date format, please specify", exact=False)).to_be_visible()
    # Finish isn't blocked by ambiguity alone - a default reading is applied.
    expect(pressable(standardise.finish_button)).not_to_have_attribute("aria-disabled", "true")

    standardise.select_month_first()
    standardise.finish()

    expect(datasets_page.status_badge()).to_have_text("Complete")

    download = datasets_page.download_standardised("qa_ambiguous_date")
    path = tmp_path / "ambiguous_date.xlsx"
    download.save_as(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = list(wb["RawData"].iter_rows(values_only=True))
    # Month-first: 03/04/2024 -> Mar 4, 05/06/2024 -> May 6.
    assert rows[1:] == [("2024-03-04",), ("2024-05-06",)]


@pytest.mark.regression
def test_unreliable_date_column_blocks_finish(datasets_page: DatasetsPage, tmp_path):
    # Both rows parse as *some* date individually (so the column still
    # defaults to Date type), but no single format string parses both -
    # "03/04/2024" is a slash D/M/Y-or-M/D/Y style, "2024-05-06" is ISO.
    path = write_file(tmp_path, "qa_unreliable_date.csv", b"Col\n03/04/2024\n2024-05-06\n")
    datasets_page.upload(path)

    standardise = open_standardise(datasets_page, "qa_unreliable_date")
    standardise.go_to_page_2()

    expect(standardise.unreliable_date_error("Col")).to_be_visible()
    expect(pressable(standardise.finish_button)).to_have_attribute("aria-disabled", "true")
